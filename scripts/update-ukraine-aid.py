#!/usr/bin/env python3
"""
Met à jour src/lib/data/ukraine-aid.json depuis le Ukraine Support Tracker de l'Institut de
Kiel (LA référence mondiale qui consolide l'aide à l'Ukraine par pays et par type).

Le workflow lance ce script régulièrement ; il télécharge la DERNIÈRE release publiée par
l'Institut (le lien Excel change à chaque release, mais la page pointe toujours vers la plus
récente), parse la feuille « Country Summary (€) » et régénère le JSON. Le workflow committe
le fichier s'il a changé → le site se redéploie avec les chiffres à jour.

Dépendances : requests, openpyxl.
"""
import re, json, sys, datetime, io
import requests
import openpyxl

PAGE = "https://www.ifw-kiel.de/topics/war-against-ukraine/ukraine-support-tracker/"
BASE = "https://www.ifw-kiel.de"
OUT = "src/lib/data/ukraine-aid.json"
UA = {"User-Agent": "Mozilla/5.0 (compatible; lapolitiquecestsimple-bot/1.0)"}

FR = {'United States': 'États-Unis', 'United Kingdom': 'Royaume-Uni', 'Germany': 'Allemagne', 'Canada': 'Canada', 'Japan': 'Japon', 'Denmark': 'Danemark', 'Sweden': 'Suède', 'Norway': 'Norvège', 'Netherlands': 'Pays-Bas', 'France': 'France', 'Poland': 'Pologne', 'Italy': 'Italie', 'Finland': 'Finlande', 'Spain': 'Espagne', 'Belgium': 'Belgique', 'Austria': 'Autriche', 'Australia': 'Australie', 'Czech Republic': 'Tchéquie', 'Czechia': 'Tchéquie', 'Estonia': 'Estonie', 'Lithuania': 'Lituanie', 'Latvia': 'Lettonie', 'Switzerland': 'Suisse', 'Portugal': 'Portugal', 'Greece': 'Grèce', 'Ireland': 'Irlande', 'South Korea': 'Corée du Sud', 'New Zealand': 'Nouvelle-Zélande', 'Iceland': 'Islande', 'Luxembourg': 'Luxembourg', 'Slovakia': 'Slovaquie', 'Slovenia': 'Slovénie', 'Croatia': 'Croatie', 'Romania': 'Roumanie', 'Bulgaria': 'Bulgarie', 'Hungary': 'Hongrie', 'Turkey': 'Turquie', 'Taiwan': 'Taïwan'}


def latest_xlsx_url():
    html = requests.get(PAGE, headers=UA, timeout=60).text
    links = re.findall(r'href="([^"]*Ukraine_Support_Tracker_Release_[^"]*\.xlsx)"', html)
    if not links:
        raise RuntimeError("Aucun lien Excel trouvé sur la page Kiel")
    def relnum(u):
        m = re.search(r'Release_(\d+)', u)
        return int(m.group(1)) if m else 0
    best = sorted(links, key=relnum)[-1]
    if best.startswith("/"):
        best = BASE + best
    return best, relnum(best)


def parse(xlsx_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True, read_only=True)
    name = [n for n in wb.sheetnames if n.startswith('Country Summary') and '$' not in n][0]
    rows = list(wb[name].iter_rows(values_only=True))
    countries, eu, total_all = [], None, None
    rnd = lambda v: round(v, 2) if isinstance(v, (int, float)) else 0
    for r in rows[9:]:
        c = r[1]
        if not c or not isinstance(c, str) or not isinstance(r[7], (int, float)):
            continue
        rec = {'name': FR.get(c, c), 'en': c, 'financial': rnd(r[4]), 'humanitarian': rnd(r[5]), 'military': rnd(r[6]), 'total': rnd(r[7]), 'euMember': bool(r[2])}
        if c == 'Total':
            total_all = rec['total']; continue
        if c.startswith('EU (') or c == 'European Union':
            eu = rec; continue
        countries.append(rec)
    countries.sort(key=lambda x: -x['total'])
    return countries, eu, total_all


def main():
    url, rel = latest_xlsx_url()
    print(f"Release {rel} -> {url}")
    xlsx = requests.get(url, headers=UA, timeout=120).content
    countries, eu, total_all = parse(xlsx)
    fr = next(c for c in countries if c['en'] == 'France')
    fr_rank = countries.index(fr) + 1
    top = countries[:12]
    if fr not in top:
        top = top + [fr]
    donors = [{'name': c['name'], 'total': c['total'], 'military': c['military'], 'financial': c['financial'], 'humanitarian': c['humanitarian'], 'isFrance': c['en'] == 'France'} for c in top]
    if eu:
        donors.append({'name': 'Union européenne', 'total': eu['total'], 'military': eu['military'], 'financial': eu['financial'], 'humanitarian': eu['humanitarian'], 'isEU': True})
    donors.sort(key=lambda x: -x['total'])
    out = {
        'source': 'Ukraine Support Tracker — Institut de Kiel (IfW)',
        'sourceUrl': PAGE,
        'release': f'Release {rel}',
        'unit': 'Md€',
        'metric': 'Aide bilatérale allouée, cumul depuis janvier 2022',
        'updatedAt': datetime.datetime.now(datetime.timezone.utc).strftime('%Y-%m-%d'),
        'france': {'financial': fr['financial'], 'humanitarian': fr['humanitarian'], 'military': fr['military'], 'total': fr['total'], 'rank': fr_rank},
        'eu': eu, 'totalAll': total_all, 'donors': donors, 'nbCountries': len(countries),
    }
    with open(OUT, 'w', encoding='utf8') as f:
        f.write(json.dumps(out, ensure_ascii=False, indent=2) + "\n")
    print(f"[ok] {OUT} — France {fr['total']} Md€ (rang {fr_rank}), {len(countries)} pays")


if __name__ == '__main__':
    try:
        main()
    except Exception as e:
        print(f"Erreur : {e}", file=sys.stderr)
        sys.exit(1)
